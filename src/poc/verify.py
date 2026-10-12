"""Independent offline verification of rendered decision briefs."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from collections import Counter, defaultdict, deque
from io import BytesIO
from pathlib import Path
from typing import Any

import yaml
from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table as WordTable
from PIL import Image

from poc.render_docx import ROOT, _normalized

TOKEN_RE = re.compile(r"[\w]+(?:['’][\w]+)*", re.UNICODE)
PRIVATE_USE = range(0xE000, 0xF900)


def _tokens(text: str) -> list[str]:
    return [token.casefold() for token in TOKEN_RE.findall(text)]


def _clean_text(text: str) -> str:
    return "".join(char for char in text if ord(char) not in PRIVATE_USE)


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _source_text(element: dict[str, Any]) -> str:
    kind = element["type"]
    if kind in {"paragraph", "bullet"}:
        return _clean_text("".join(span["text"] for span in element.get("spans", [])))
    if kind == "heading":
        return " ".join(part for part in (element.get("number"), element.get("title")) if part)
    if kind == "table":
        return " ".join(cell for row in element.get("rows", []) for cell in row)
    if kind == "page_number":
        return element.get("value", "")
    return ""


def _source_words(payload: dict[str, Any]) -> tuple[Counter[str], dict[str, Counter[str]]]:
    total: Counter[str] = Counter()
    by_id: dict[str, Counter[str]] = {}
    for element in payload["elements"]:
        if element.get("disposition") == "regenerated" or element["type"] in {"figure", "page_number"}:
            continue
        words = Counter(_tokens(_source_text(element)))
        by_id[element["id"]] = words
        total.update(words)
    return total, by_id


def _iter_body_blocks(document: Any):
    """Yield top-level paragraphs and tables in document order."""
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield WordTable(child, document)


def _paragraph_images(paragraph: Any) -> list[bytes]:
    blobs = []
    for blip in paragraph._p.xpath(".//a:blip"):
        relation_id = blip.get(qn("r:embed"))
        if relation_id:
            blobs.append(paragraph.part.related_parts[relation_id].blob)
    return blobs


def _all_output_text(document: Any) -> list[tuple[str, str]]:
    located: list[tuple[str, str]] = []
    for index, block in enumerate(_iter_body_blocks(document), start=1):
        if isinstance(block, Paragraph):
            if block.text:
                style = block.style.name if block.style else "paragraph"
                located.append((block.text, f"body paragraph {index} ({style})"))
        else:
            for row_idx, row in enumerate(block.rows, start=1):
                for col_idx, cell in enumerate(row.cells, start=1):
                    if cell.text:
                        located.append((cell.text, f"table block {index}, row {row_idx}, column {col_idx}"))
    for section_index, section in enumerate(document.sections, start=1):
        for label, part in (("header", section.header), ("footer", section.footer)):
            for paragraph_index, paragraph in enumerate(part.paragraphs, start=1):
                if paragraph.text:
                    located.append((paragraph.text, f"section {section_index} {label} paragraph {paragraph_index}"))
    return located


def _generated_allowance(payload: dict[str, Any], config: dict[str, Any]) -> Counter[str]:
    """Build a token allowance from renderer-generated text, not observed arbitrary output."""
    texts: list[str] = ["Contents", "Page", "Structure inside images may be lost in transcription."]
    section_titles = [section["title"] for section in config["sections"]]
    texts.extend(section_titles)  # Static contents list.
    texts.extend(section_titles)  # Target section headings.
    texts.extend(item["title"] for item in config["appendices"])
    texts.extend([
        "AI-generated figure id", "AI-generated source page", "AI-generated hash prefix",
        "AI-generated caption", "AI-generated transcribed text", "AI-generated source part hashes",
        "AI-generated description:",
    ])
    figures = [element for element in payload["elements"] if element["type"] == "figure"]
    for number, figure in enumerate(figures, start=1):
        texts.append(f"Figure {number}")
        ai = figure.get("ai") or {}
        caption = ai.get("caption", "[missing enrichment result]")
        transcript = " | ".join(ai.get("transcribed_text", [])) or "[none]"
        part_hashes = [part.get("sha256", "") for part in figure.get("parts", [])]
        texts.extend([
            f"AI-generated description: {caption}",
            figure["id"], str(figure["page"]), figure.get("sha256", "")[:12], caption,
            transcript, ", ".join(part_hashes) or figure.get("sha256", ""),
        ])
    classifications = json.loads((ROOT / "data/work" / payload["doc_id"] / "classification.json").read_text())
    if any(item.get("canonical_type") is None for item in classifications.get("items", [])):
        texts.append("Unclassified Source Content")
    section_by_type = {canonical: section for section in config["sections"] for canonical in section["accepts"]}
    for item in classifications["items"]:
        title = item["title"]
        canonical = item.get("canonical_type")
        target = section_by_type[canonical]["title"] if canonical in section_by_type else "Unclassified Source Content"
        texts.append(f"{item['heading_id']}: source number {item.get('number') or '[none]'}; title {title}; target {target}")
    texts.extend(str(number) for number in range(1, payload["page_count"] + 1))
    allowance: Counter[str] = Counter()
    for text in texts:
        allowance.update(_tokens(text))
    return allowance


def _text_metrics(payload: dict[str, Any], output_text: list[tuple[str, str]], config: dict[str, Any]) -> dict[str, Any]:
    source, by_id = _source_words(payload)
    output: Counter[str] = Counter()
    output_locations: dict[str, list[str]] = defaultdict(list)
    for text, location in output_text:
        tokens = _tokens(text)
        output.update(tokens)
        output_locations.update({token: [*output_locations[token], location] for token in set(tokens)})
    missing_counter = source - output
    allowance = _generated_allowance(payload, config)
    excess_counter = output - source
    allowed_used = Counter({token: min(count, allowance[token]) for token, count in excess_counter.items()})
    excess_counter.subtract(allowed_used)
    excess_counter += Counter()

    missing: list[dict[str, Any]] = []
    remaining_output = output.copy()
    for element in payload["elements"]:
        element_id = element["id"]
        words = by_id.get(element_id)
        if not words:
            continue
        for word, count in words.items():
            available = remaining_output[word]
            lost = max(0, count - available)
            if lost:
                missing.append({"element_id": element_id, "word": word, "count": lost})
            remaining_output[word] = max(0, available - count)
    excess: list[dict[str, Any]] = []
    for word, count in excess_counter.items():
        locations = output_locations.get(word, ["unknown output location"])
        for index in range(count):
            excess.append({"word": word, "location": locations[min(index, len(locations) - 1)], "possible_fabrication": True})
    source_count = sum(source.values())
    output_source_hits = sum((source & output).values())
    return {
        "source_word_count": source_count,
        "output_word_count": sum(output.values()),
        "recalled_source_word_count": output_source_hits,
        "word_recall": output_source_hits / source_count if source_count else None,
        "word_recall_reason": None if source_count else "No source word tokens remain after regenerated elements are excluded.",
        "missing_word_count": sum(missing_counter.values()),
        "missing_words": missing,
        "excess_word_count": sum(excess_counter.values()),
        "excess_words": excess,
    }


def _norm_cell(text: str) -> str:
    return " ".join(_clean_text(text).split())


def _table_signature(rows: list[list[str]]) -> tuple[tuple[str, ...], ...]:
    return tuple(tuple(_norm_cell(cell) for cell in row) for row in rows)


def _source_tables(payload: dict[str, Any]) -> list[dict[str, Any]]:
    return [e for e in payload["elements"] if e["type"] == "table" and e.get("role") != "toc"]


def _output_tables(document: Any) -> list[dict[str, Any]]:
    output = []
    for index, table in enumerate(document.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        first = rows[0] if rows else []
        if first and all("AI-generated" in cell for cell in first):
            continue  # Appendix 1 generated index table.
        output.append({"index": index, "rows": rows})
    return output


def _table_metrics(payload: dict[str, Any], document: Any) -> dict[str, Any]:
    expected = _source_tables(payload)
    actual = _output_tables(document)
    defects: list[dict[str, Any]] = []
    comparisons = []
    if len(expected) != len(actual):
        defects.append({"code": "table_count_mismatch", "message": f"Expected {len(expected)} source tables, found {len(actual)} output tables."})
    for index in range(max(len(expected), len(actual))):
        source_table = expected[index] if index < len(expected) else None
        output_table = actual[index] if index < len(actual) else None
        source_rows = source_table["rows"] if source_table else []
        output_rows = output_table["rows"] if output_table else []
        shape_expected = [len(source_rows), max((len(row) for row in source_rows), default=0)]
        shape_actual = [len(output_rows), max((len(row) for row in output_rows), default=0)]
        cell_equal = _table_signature(source_rows) == _table_signature(output_rows)
        record = {
            "source_element_id": source_table["id"] if source_table else None,
            "output_table_index": output_table["index"] if output_table else None,
            "expected_shape": shape_expected,
            "actual_shape": shape_actual,
            "shape_equal": shape_expected == shape_actual,
            "cells_equal_after_whitespace_normalisation": cell_equal,
        }
        comparisons.append(record)
        if shape_expected != shape_actual:
            defects.append({"code": "table_shape_mismatch", "element_id": record["source_element_id"], "message": f"Table shape {shape_actual} differs from expected {shape_expected}."})
        elif not cell_equal:
            defects.append({"code": "table_cell_mismatch", "element_id": record["source_element_id"], "message": "One or more table cells differ after whitespace normalisation."})
    return {"expected_count": len(expected), "actual_count": len(actual), "matching_count": sum(item["shape_equal"] and item["cells_equal_after_whitespace_normalisation"] for item in comparisons), "all_equal": not defects, "comparisons": comparisons, "defects": defects}


def _heading_metrics(payload: dict[str, Any], document: Any) -> dict[str, Any]:
    source_headings = [e for e in payload["elements"] if e["type"] == "heading"]
    output_headings: list[tuple[int, str, str]] = []
    for paragraph in document.paragraphs:
        style = paragraph.style.name if paragraph.style else ""
        match = re.match(r"Heading\s+(\d+)", style)
        if match:
            output_headings.append((int(match.group(1)), paragraph.text, style))
    available = Counter((level, title) for level, title, _ in output_headings)
    defects = []
    matches = []
    for element in source_headings:
        expected_level = int(element.get("level", 1)) + 1  # One level below the generated target section heading.
        signature = (expected_level, element["title"])
        if available[signature] > 0:
            available[signature] -= 1
            matches.append({"element_id": element["id"], "title": element["title"], "source_level": element["level"], "output_heading_level": expected_level, "matched": True})
        else:
            matches.append({"element_id": element["id"], "title": element["title"], "source_level": element["level"], "output_heading_level": expected_level, "matched": False})
            defects.append({"code": "heading_missing_or_wrong_level", "element_id": element["id"], "message": f"Heading {element['title']!r} does not appear verbatim at output Heading {expected_level}."})
    return {"source_count": len(source_headings), "matched_count": sum(item["matched"] for item in matches), "matches": matches, "output_headings": [{"level": level, "title": title} for level, title, _ in output_headings], "defects": defects}


def _expected_target_figures(payload: dict[str, Any], placements: dict[str, Any], config: dict[str, Any]) -> list[dict[str, Any]]:
    section_order = {section["id"]: index for index, section in enumerate(config["sections"])}
    figures = [element for element in payload["elements"] if element["type"] == "figure"]
    def key(element: dict[str, Any]) -> tuple[int, int]:
        target = placements.get(element["id"], {}).get("target_section", "cover")
        return (-1 if target == "cover" else section_order.get(target, 999), payload["elements"].index(element))
    return sorted(figures, key=key)


def _docx_images(docx_path: Path, document: Any) -> tuple[list[dict[str, Any]], Counter[str]]:
    with zipfile.ZipFile(docx_path) as archive:
        media_names = [name for name in archive.namelist() if name.startswith("word/media/")]
        media_hashes = Counter(_sha(archive.read(name)) for name in media_names)
    inline = []
    for block in _iter_body_blocks(document):
        if isinstance(block, Paragraph):
            for blob in _paragraph_images(block):
                inline.append({"sha256": _sha(blob), "location": f"body paragraph image {len(inline) + 1}"})
    return inline, media_hashes


def _rebuild_composite(doc_dir: Path, figure: dict[str, Any], tolerance: int) -> bytes:
    parts = figure["parts"]
    reference = parts[0]
    x0, y0, _, _ = figure["bbox"]
    scale_x = reference["width"] / (reference["bbox"][2] - reference["bbox"][0])
    scale_y = reference["height"] / (reference["bbox"][3] - reference["bbox"][1])
    decoded: list[tuple[Image.Image, int, int]] = []
    for part in parts:
        with Image.open(doc_dir / part["image_path"]) as source:
            source.load()
            image = source.convert("RGBA")
        mask_path = part.get("mask_path")
        if mask_path:
            with Image.open(doc_dir / mask_path) as mask_source:
                mask_source.load()
                mask = mask_source.convert("L")
            if mask.size != image.size:
                mask = mask.resize(image.size, Image.Resampling.NEAREST)
            image.putalpha(mask)
        expected = (
            max(1, round((part["bbox"][2] - part["bbox"][0]) * scale_x)),
            max(1, round((part["bbox"][3] - part["bbox"][1]) * scale_y)),
        )
        if max(abs(expected[0] - image.width), abs(expected[1] - image.height)) > tolerance:
            image = image.resize(expected, Image.Resampling.LANCZOS)
        x = round((part["bbox"][0] - x0) * scale_x)
        y = round((part["bbox"][1] - y0) * scale_y)
        decoded.append((image, x, y))
    width = max(1, max(x + image.width for image, x, _ in decoded))
    height = max(1, max(y + image.height for image, _, y in decoded))
    canvas = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    for image, x, y in decoded:
        canvas.alpha_composite(image, dest=(x, y))
    buffer = BytesIO()
    canvas.save(buffer, format="PNG", optimize=False)
    return buffer.getvalue()


def _image_metrics(doc_id: str, payload: dict[str, Any], placements: dict[str, Any], config: dict[str, Any], docx_path: Path, document: Any) -> dict[str, Any]:
    figures = [element for element in payload["elements"] if element["type"] == "figure"]
    expected_hashes = Counter(element["sha256"] for element in figures)
    inline, media_hashes = _docx_images(docx_path, document)
    actual_inline_hashes = Counter(item["sha256"] for item in inline)
    defects = []
    if len(media_hashes) != len(figures) or set(media_hashes) != set(expected_hashes):
        media_defects = sorted(element["id"] for element in figures if media_hashes[element["sha256"]] != expected_hashes[element["sha256"]])
        defects.append({"code": "image_set_or_count_mismatch", "element_ids": media_defects, "message": f"DOCX media has {len(media_hashes)} files; source has {len(figures)} display images."})
    if actual_inline_hashes != expected_hashes:
        defects.append({"code": "image_relationship_mismatch", "message": "Inline image relationships do not contain the source display-image hashes with matching counts."})
    expected_order = _expected_target_figures(payload, placements, config)
    expected_sequence = [figure["sha256"] for figure in expected_order]
    actual_sequence = [item["sha256"] for item in inline]
    if expected_sequence != actual_sequence:
        mismatched_ids = [expected_order[index]["id"] for index in range(min(len(expected_sequence), len(actual_sequence))) if expected_sequence[index] != actual_sequence[index]]
        defects.append({"code": "image_placement_mismatch", "element_ids": mismatched_ids, "message": "The image hashes appear in a different figure order than the source-to-target placement order."})
    extraction_cfg = yaml.safe_load((ROOT / "config/extraction.yaml").read_text())
    tolerance = int(extraction_cfg["images"]["native_rounding_tolerance_px"])
    composites = []
    doc_dir = ROOT / "data/work" / doc_id
    for figure in figures:
        if len(figure.get("parts", [])) <= 1:
            continue
        part_hashes = []
        for part in figure["parts"]:
            part_hash = _sha((doc_dir / part["image_path"]).read_bytes())
            part_hashes.append({"part_id": part["id"], "expected": part["sha256"], "actual": part_hash, "matches": part_hash == part["sha256"]})
        parts_match = all(item["matches"] for item in part_hashes)
        rebuilt = _sha(_rebuild_composite(doc_dir, figure, tolerance))
        matches = rebuilt == figure["sha256"] and parts_match
        composites.append({"figure_id": figure["id"], "part_hashes": part_hashes, "parts_match": parts_match, "rebuilt_sha256": rebuilt, "display_sha256": figure["sha256"], "matches": matches})
        if not parts_match:
            defects.append({"code": "source_part_hash_mismatch", "element_id": figure["id"], "message": "At least one composite part does not match its recorded SHA-256."})
        if not matches:
            defects.append({"code": "composite_rebuild_mismatch", "element_id": figure["id"], "message": "Composite rebuilt from exact stored parts does not match the stored display image."})
    return {
        "source_display_image_count": len(figures),
        "docx_media_file_count": len(media_hashes),
        "docx_inline_image_count": len(inline),
        "source_hashes": dict(expected_hashes),
        "media_hashes": dict(media_hashes),
        "set_equal": set(media_hashes) == set(expected_hashes),
        "count_equal": len(media_hashes) == len(figures),
        "inline_relationships_equal": actual_inline_hashes == expected_hashes,
        "placement_order_equal": expected_sequence == actual_sequence,
        "composites": composites,
        "defects": defects,
    }


def _mapping_metrics(doc_id: str, answer_key: dict[str, Any]) -> dict[str, Any]:
    classifications = json.loads((ROOT / "data/work" / doc_id / "classification.json").read_text())
    expected = answer_key.get(doc_id, {}).get("section_types", {})
    actual = {item["number"]: item["canonical_type"] for item in classifications["items"] if "." not in item["number"]}
    heading_ids = {item["number"]: item["heading_id"] for item in classifications["items"] if "." not in item["number"]}
    comparisons = []
    for number, expected_type in expected.items():
        actual_type = actual.get(number)
        comparisons.append({"element_id": heading_ids.get(number), "number": number, "expected": expected_type, "actual": actual_type, "correct": actual_type == expected_type})
    correct = sum(item["correct"] for item in comparisons)
    return {"correct": correct, "total": len(comparisons), "accuracy": correct / len(comparisons) if comparisons else None, "comparisons": comparisons}


def _normalise_token(text: str) -> str:
    return re.sub(r"[-\s]+", "", text.casefold())


def _image_token_metrics(doc_id: str, payload: dict[str, Any], answer_key: dict[str, Any]) -> dict[str, Any]:
    expected_groups = list(answer_key.get(doc_id, {}).get("image_tokens", {}).items())
    figures = [element for element in payload["elements"] if element["type"] == "figure"]
    per_figure = []
    for index, (key_name, expected_tokens) in enumerate(expected_groups):
        figure = figures[index] if index < len(figures) else None
        ai = figure.get("ai") if figure else None
        if ai is None:
            per_figure.append({"figure_id": figure["id"] if figure else None, "key_group": key_name, "strict": "NOT MEASURED", "normalised": "NOT MEASURED", "reason": "Cached AI enrichment result is missing."})
            continue
        if not expected_tokens:
            per_figure.append({"figure_id": figure["id"], "page": figure["page"], "key_group": key_name, "strict": "NOT MEASURED", "normalised": "NOT MEASURED", "reason": "The answer key selects no image keywords for this figure."})
            continue
        transcript = " ".join(ai.get("transcribed_text", []))
        strict_matches = [token for token in expected_tokens if token.casefold() in transcript.casefold()]
        normalized_text = _normalise_token(transcript)
        normalized_matches = [token for token in expected_tokens if _normalise_token(token) in normalized_text]
        denominator = len(expected_tokens)
        per_figure.append({
            "figure_id": figure["id"], "page": figure["page"], "key_group": key_name,
            "strict": {"matched": len(strict_matches), "total": denominator, "recall": len(strict_matches) / denominator if denominator else None, "matched_tokens": strict_matches, "missing_tokens": [token for token in expected_tokens if token not in strict_matches]},
            "normalised": {"matched": len(normalized_matches), "total": denominator, "recall": len(normalized_matches) / denominator if denominator else None, "matched_tokens": normalized_matches, "missing_tokens": [token for token in expected_tokens if token not in normalized_matches]},
            "label": "indicative; measures keyword presence only",
        })
    measured = [item for item in per_figure if isinstance(item.get("strict"), dict)]
    strict_total = sum(item["strict"]["total"] for item in measured)
    strict_matched = sum(item["strict"]["matched"] for item in measured)
    norm_total = sum(item["normalised"]["total"] for item in measured)
    norm_matched = sum(item["normalised"]["matched"] for item in measured)
    return {
        "label": "indicative; measures keyword presence only",
        "strict": {"matched": strict_matched, "total": strict_total, "recall": strict_matched / strict_total if strict_total else None},
        "normalised": {"matched": norm_matched, "total": norm_total, "recall": norm_matched / norm_total if norm_total else None},
        "per_figure": per_figure,
    }


def _section_order_metrics(doc_id: str, payload: dict[str, Any], document: Any, placements: dict[str, Any], config: dict[str, Any]) -> dict[str, Any]:
    target_titles = {section["title"]: section["id"] for section in config["sections"]}
    actual_order = []
    current_section: str | None = None
    actual_by_section: dict[str, list[tuple[Any, str]]] = defaultdict(list)
    section_wrong: list[dict[str, Any]] = []
    for block in _iter_body_blocks(document):
        if isinstance(block, Paragraph):
            style = block.style.name if block.style else ""
            level = int(style.split()[-1]) if style.startswith("Heading ") and style.split()[-1].isdigit() else None
            text = block.text
            if level == 1:
                current_section = target_titles.get(text)
                if current_section:
                    actual_order.append(current_section)
                else:
                    current_section = None
                continue
            if current_section is None:
                continue
            images = _paragraph_images(block)
            if images:
                for blob in images:
                    actual_by_section[current_section].append(("figure", _sha(blob)))
            elif level in {2, 3}:
                actual_by_section[current_section].append(("heading", level, _normalized(text)))
            elif text and not text.startswith(("Figure ", "AI-generated description:")):
                actual_by_section[current_section].append(("text", _normalized(_clean_text(text))))
        else:
            if current_section is not None:
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                if rows and all("AI-generated" in cell for cell in rows[0]):
                    continue
                actual_by_section[current_section].append(("table", _table_signature(rows)))
    expected_order = [section["id"] for section in config["sections"]]
    order_defects = []
    if actual_order != expected_order:
        order_defects.append({"code": "target_section_order", "message": f"Target sections appear as {actual_order}; expected {expected_order}."})

    expected_by_section: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    owner_by_signature: dict[tuple[Any, ...], set[str]] = defaultdict(set)
    ids_by_section_signature: dict[str, dict[tuple[Any, ...], list[str]]] = defaultdict(lambda: defaultdict(list))
    for element in payload["elements"]:
        if element["type"] in {"page_number", "figure"} and element["type"] == "page_number":
            continue
        location = placements.get(element["id"], {})
        section = location.get("target_section")
        if section not in target_titles.values() or element.get("disposition") != "placed":
            continue
        kind = element["type"]
        if kind == "heading":
            signature = ("heading", int(element.get("level", 1)) + 1, _normalized(element["title"]))
        elif kind in {"paragraph", "bullet"}:
            signature = ("text", _normalized(_clean_text(_source_text(element))))
        elif kind == "table":
            signature = ("table", _table_signature(element["rows"]))
        elif kind == "figure":
            signature = ("figure", element["sha256"])
        else:
            continue
        expected_by_section[section].append(signature)
        owner_by_signature[signature].add(section)
        ids_by_section_signature[section][signature].append(element["id"])

    violations = []
    for section_id in target_titles.values():
        actual = actual_by_section.get(section_id, [])
        expected = expected_by_section.get(section_id, [])
        expected_counts = Counter(expected)
        sequence = []
        seen_counts: Counter[Any] = Counter()
        for signature in actual:
            if signature in owner_by_signature and section_id not in owner_by_signature[signature]:
                violations.append({"section": section_id, "element_ids": sorted(element_id for owner in owner_by_signature[signature] for element_id in ids_by_section_signature[owner][signature]), "signature": repr(signature), "message": "Source element appears under the wrong target section."})
            if expected_counts[signature] > seen_counts[signature]:
                sequence.append(expected.index(signature, sum(1 for prior in sequence if prior == signature)))
                seen_counts[signature] += 1
        if any(left > right for left, right in zip(sequence, sequence[1:])):
            violations.append({"section": section_id, "element_ids": [element_id for signatures in ids_by_section_signature[section_id].values() for element_id in signatures], "message": "Source elements are out of relative order."})
    if violations:
        order_defects.append({"code": "source_order_violation", "message": "One or more source elements moved to a different target section or changed relative order.", "violations": violations})
    return {"expected_section_order": expected_order, "actual_section_order": actual_order, "section_order_preserved": actual_order == expected_order, "relative_order_preserved": not violations, "violations": violations, "defects": order_defects}


def _contents_metrics(payload: dict[str, Any], heading_metrics: dict[str, Any]) -> dict[str, Any]:
    topics = []
    for element in payload["elements"]:
        if element["type"] == "table" and element.get("role") == "toc":
            for row in element["rows"][1:]:
                if row and row[0].strip():
                    topics.append({"topic": row[0].strip(), "element_id": element["id"]})
    output_titles = [item["title"] for item in heading_metrics["output_headings"]]
    missing = [topic for topic in topics if topic["topic"] not in output_titles]
    defects = [{"code": "contents_topic_missing", "element_id": topic["element_id"], "topic": topic["topic"], "message": f"Contents topic {topic['topic']!r} is not present as an output heading."} for topic in missing]
    return {"topic_count": len(topics), "matched_count": len(topics) - len(missing), "missing_topics": missing, "defects": defects}


def _ledger_metrics(payload: dict[str, Any], placement: dict[str, Any]) -> dict[str, Any]:
    elements = payload["elements"]
    ids = [element["id"] for element in elements]
    duplicates = sorted({element_id for element_id, count in Counter(ids).items() if count != 1})
    placements = placement.get("placements", {})
    missing_placement = sorted(set(ids) - set(placements))
    unknown_placement = sorted(set(placements) - set(ids))
    expected_regenerated = {
        element["id"] for element in elements
        if element["type"] == "page_number"
        or element["type"] == "table" and element.get("role") == "toc"
        or element["type"] == "paragraph" and _normalized(_source_text(element)).casefold() == "contents"
    }
    actual_regenerated = {element["id"] for element in elements if element.get("disposition") == "regenerated"}
    counts = Counter(element.get("disposition", "missing") for element in elements)
    defects = []
    if duplicates:
        defects.append({"code": "duplicate_element_id", "ids": duplicates, "message": "Source element IDs are not unique."})
    if missing_placement or unknown_placement or len(placements) != len(ids):
        defects.append({"code": "placement_ledger_mismatch", "missing_ids": missing_placement, "unknown_ids": unknown_placement, "message": "Placement ledger does not contain exactly one mapping for every source element."})
    if any(element.get("disposition") not in {"placed", "regenerated", "unplaced"} for element in elements):
        defects.append({"code": "pending_disposition", "message": "At least one source element has a missing or pending disposition."})
    if counts["unplaced"] != 0:
        defects.append({"code": "unplaced_elements", "count": counts["unplaced"], "message": "Unplaced source elements must be zero."})
    if actual_regenerated != expected_regenerated:
        defects.append({"code": "regenerated_set_mismatch", "unexpected": sorted(actual_regenerated - expected_regenerated), "missing": sorted(expected_regenerated - actual_regenerated), "message": "Regenerated elements are not exactly page numbers, the contents table, and the Contents title."})
    return {"element_count": len(elements), "counts": {key: counts[key] for key in ("placed", "regenerated", "unplaced", "pending")}, "expected_regenerated_ids": sorted(expected_regenerated), "actual_regenerated_ids": sorted(actual_regenerated), "placement_count": len(placements), "defects": defects}


def _private_use_metrics(output_text: list[tuple[str, str]]) -> dict[str, Any]:
    found = []
    for text, location in output_text:
        for char in text:
            codepoint = ord(char)
            if codepoint in PRIVATE_USE:
                found.append({"character": f"U+{codepoint:04X}", "location": location})
    return {"count": len(found), "characters": found, "defects": [{"code": "private_use_character", **item, "message": f"Private-use character {item['character']} appears in output text."} for item in found]}


def verify_document(
    doc_id: str,
    docx_path: Path | None = None,
    *,
    answer_key_override: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Verify one produced DOCX against its independent source IR and local key."""
    config = yaml.safe_load((ROOT / "config/target_format.yaml").read_text())
    answer_key = answer_key_override or json.loads((ROOT / "docs/answer_key.json").read_text())
    doc_dir = ROOT / "data/work" / doc_id
    payload = json.loads((doc_dir / "elements.json").read_text())
    placement = json.loads((doc_dir / "placement.json").read_text())
    if docx_path is None:
        docx_path = ROOT / "data/output" / f"{doc_id}_decision_brief.docx"
    docx_path = Path(docx_path)
    document = WordDocument(docx_path)
    output_text = _all_output_text(document)
    ledger = _ledger_metrics(payload, placement)
    text = _text_metrics(payload, output_text, config)
    tables = _table_metrics(payload, document)
    headings = _heading_metrics(payload, document)
    images = _image_metrics(doc_id, payload, placement["placements"], config, docx_path, document)
    order = _section_order_metrics(doc_id, payload, document, placement["placements"], config)
    contents = _contents_metrics(payload, headings)
    mapping = _mapping_metrics(doc_id, answer_key)
    tokens = _image_token_metrics(doc_id, payload, answer_key)
    private_use = _private_use_metrics(output_text)
    categories = {
        "ledger": ledger, "text": text, "tables": tables, "headings": headings,
        "images": images, "order": order, "contents": contents,
        "mapping": mapping, "image_transcription": tokens, "private_use": private_use,
    }
    defects = []
    for category, details in categories.items():
        for defect in details.get("defects", []):
            defects.append({"category": category, **defect})
    for item in text["missing_words"]:
        defects.append({"category": "text", "code": "text_missing", **item, "message": f"Source word {item['word']!r} is missing from output."})
    for item in text["excess_words"]:
        defects.append({"category": "text", "code": "text_excess", **item, "message": f"Output word {item['word']!r} is not source or allowed generated text; possible fabrication."})
    for item in mapping["comparisons"]:
        if not item["correct"]:
            defects.append({"category": "mapping", "code": "mapping_mismatch", **item, "message": "Section classification differs from the answer key."})
    return {"doc_id": doc_id, "docx_file": str(docx_path), **categories, "defects": defects}


def _aggregate(results: list[dict[str, Any]]) -> dict[str, Any]:
    docs = [result["metrics"] for result in results]
    def total(path: tuple[str, ...]) -> int:
        value = 0
        for doc in docs:
            current: Any = doc
            for part in path:
                current = current[part]
            value += int(current or 0)
        return value
    source_words = total(("text", "source_word_count"))
    recalled = total(("text", "recalled_source_word_count"))
    strict_match = total(("image_transcription", "strict", "matched"))
    strict_total = total(("image_transcription", "strict", "total"))
    normalized_match = total(("image_transcription", "normalised", "matched"))
    normalized_total = total(("image_transcription", "normalised", "total"))
    map_correct = total(("mapping", "correct"))
    map_total = total(("mapping", "total"))
    headings_matched = total(("headings", "matched_count"))
    headings_total = total(("headings", "source_count"))
    contents_matched = total(("contents", "matched_count"))
    contents_total = total(("contents", "topic_count"))
    return {
        "documents": len(results),
        "text": {"source_word_count": source_words, "recalled_source_word_count": recalled, "word_recall": recalled / source_words if source_words else None, "missing_word_count": total(("text", "missing_word_count")), "excess_word_count": total(("text", "excess_word_count"))},
        "tables": {"expected_count": total(("tables", "expected_count")), "matching_count": total(("tables", "matching_count")), "shapes_equal": all(all(comparison["shape_equal"] for comparison in item["metrics"]["tables"]["comparisons"]) for item in results), "cells_equal": all(all(comparison["cells_equal_after_whitespace_normalisation"] for comparison in item["metrics"]["tables"]["comparisons"]) for item in results)},
        "images": {"source_display_image_count": total(("images", "source_display_image_count")), "docx_media_file_count": total(("images", "docx_media_file_count")), "composites_rebuilt": sum(len(item["metrics"]["images"]["composites"]) for item in results), "composites_matching": sum(sum(1 for composite in item["metrics"]["images"]["composites"] if composite["matches"]) for item in results)},
        "mapping": {"correct": map_correct, "total": map_total, "accuracy": map_correct / map_total if map_total else None},
        "headings": {"matched": headings_matched, "total": headings_total, "accuracy": headings_matched / headings_total if headings_total else None},
        "contents": {"matched": contents_matched, "total": contents_total, "all_topics_headings": contents_matched == contents_total},
        "order": {"documents_in_section_order": sum(item["metrics"]["order"]["section_order_preserved"] for item in results), "documents_with_relative_order": sum(item["metrics"]["order"]["relative_order_preserved"] for item in results)},
        "image_transcription": {"label": "indicative; measures keyword presence only", "strict": {"matched": strict_match, "total": strict_total, "recall": strict_match / strict_total if strict_total else None}, "normalised": {"matched": normalized_match, "total": normalized_total, "recall": normalized_match / normalized_total if normalized_total else None}},
        "ledger": {"placed": total(("ledger", "counts", "placed")), "regenerated": total(("ledger", "counts", "regenerated")), "unplaced": total(("ledger", "counts", "unplaced")), "pending": total(("ledger", "counts", "pending"))},
        "private_use_character_count": total(("private_use", "count")),
        "defect_count": sum(len(item["metrics"]["defects"]) for item in results),
    }


def _display_value(value: Any) -> str:
    if value is None:
        return "NOT MEASURED"
    if isinstance(value, float):
        return f"{value:.4%}"
    return str(value)


def _markdown_report(results: list[dict[str, Any]], overall: dict[str, Any]) -> str:
    lines = ["# Verification Report — M5", "", "All checks ran offline against local IR, cached enrichment, and generated DOCX files.", "", "## Summary", "", "| Document | Text recall | Missing words | Excess words | Tables matching | Images media/source | Mapping accuracy | Strict image keywords | Normalised image keywords | Placed / regenerated / unplaced | Defects |", "|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|"]
    for result in results:
        metrics = result["metrics"]
        lines.append(f"| {result['doc_id']} | {_display_value(metrics['text']['word_recall'])} | {metrics['text']['missing_word_count']} | {metrics['text']['excess_word_count']} | {metrics['tables']['matching_count']}/{metrics['tables']['expected_count']} | {metrics['images']['docx_media_file_count']}/{metrics['images']['source_display_image_count']} | {metrics['mapping']['correct']}/{metrics['mapping']['total']} ({_display_value(metrics['mapping']['accuracy'])}) | {metrics['image_transcription']['strict']['matched']}/{metrics['image_transcription']['strict']['total']} ({_display_value(metrics['image_transcription']['strict']['recall'])}) | {metrics['image_transcription']['normalised']['matched']}/{metrics['image_transcription']['normalised']['total']} ({_display_value(metrics['image_transcription']['normalised']['recall'])}) | {metrics['ledger']['counts']['placed']} / {metrics['ledger']['counts']['regenerated']} / {metrics['ledger']['counts']['unplaced']} | {len(metrics['defects'])} |")
    lines.extend(["", "### Overall", "", f"- Text word recall: {_display_value(overall['text']['word_recall'])}; missing {overall['text']['missing_word_count']} source words; excess {overall['text']['excess_word_count']} possible-fabrication words.", f"- Tables matching: {overall['tables']['matching_count']}/{overall['tables']['expected_count']}.", f"- Display images: {overall['images']['docx_media_file_count']}/{overall['images']['source_display_image_count']}; composites rebuilt and matching: {overall['images']['composites_matching']}/{overall['images']['composites_rebuilt']}.", f"- Mapping accuracy: {overall['mapping']['correct']}/{overall['mapping']['total']} ({_display_value(overall['mapping']['accuracy'])}).", f"- Image transcription, indicative; measures keyword presence only: strict {overall['image_transcription']['strict']['matched']}/{overall['image_transcription']['strict']['total']} ({_display_value(overall['image_transcription']['strict']['recall'])}); normalised {overall['image_transcription']['normalised']['matched']}/{overall['image_transcription']['normalised']['total']} ({_display_value(overall['image_transcription']['normalised']['recall'])}).", f"- Dispositions: {overall['ledger']['placed']} placed, {overall['ledger']['regenerated']} regenerated, {overall['ledger']['unplaced']} unplaced, {overall['ledger']['pending']} pending.", f"- Private-use characters: {overall['private_use_character_count']}; defects: {overall['defect_count']}.", "", "Image transcription token scores are indicative; they measure presence of selected keywords only."])
    for result in results:
        metrics = result["metrics"]
        lines.extend(["", f"## {result['doc_id'].title()} details", ""])
        lines.extend(["### Ledger", "", f"Disposition counts: {metrics['ledger']['counts']}. Expected regenerated set matches: {not any(d['code'] == 'regenerated_set_mismatch' for d in metrics['ledger']['defects'])}.", "", "### Text misses and possible excess", ""])
        if not metrics["text"]["missing_words"] and not metrics["text"]["excess_words"]:
            lines.append("No missing or excess words.")
        else:
            for item in metrics["text"]["missing_words"]:
                lines.append(f"- MISSING `{item['word']}` ×{item['count']} — source element `{item['element_id']}`.")
            for item in metrics["text"]["excess_words"]:
                lines.append(f"- POSSIBLE FABRICATION `{item['word']}` — {item['location']}.")
        lines.extend(["", "### Table comparison", ""])
        for comparison in metrics["tables"]["comparisons"]:
            lines.append(f"- `{comparison['source_element_id']}`: expected {comparison['expected_shape']}, output {comparison['actual_shape']}; shape {'PASS' if comparison['shape_equal'] else 'FAIL'}, cells {'PASS' if comparison['cells_equal_after_whitespace_normalisation'] else 'FAIL'}.")
        lines.extend(["", "### Heading and contents checks", ""])
        for defect in metrics["headings"]["defects"]:
            lines.append(f"- `{defect.get('element_id')}`: {defect['message']}")
        for topic in metrics["contents"]["missing_topics"]:
            lines.append(f"- `{topic['element_id']}`: missing contents topic heading `{topic['topic']}`.")
        if not metrics["headings"]["defects"] and not metrics["contents"]["missing_topics"]:
            lines.append(f"All {metrics['headings']['source_count']} source headings match at their expected Word heading levels; all {metrics['contents']['topic_count']} TOC topics are headings.")
        lines.extend(["", "### Image checks", ""])
        for defect in metrics["images"]["defects"]:
            lines.append(f"- {defect['code']}: {defect['message']}")
        for composite in metrics["images"]["composites"]:
            lines.append(f"- `{composite['figure_id']}` rebuilt composite: {'PASS' if composite['matches'] else 'FAIL'}; rebuilt SHA-256 `{composite['rebuilt_sha256']}`.")
        lines.extend(["", "### Section mapping and indicative image keywords", ""])
        for item in metrics["mapping"]["comparisons"]:
            lines.append(f"- `{item['element_id']}` section {item['number']}: expected `{item['expected']}`, extracted `{item['actual']}` — {'PASS' if item['correct'] else 'FAIL'}.")
        for item in metrics["image_transcription"]["per_figure"]:
            if not isinstance(item.get("strict"), dict):
                lines.append(f"- `{item.get('figure_id') or item['key_group']}`: NOT MEASURED — {item['reason'].rstrip('.')}.")
            else:
                strict = item["strict"]
                normalised = item["normalised"]
                lines.append(f"- `{item['figure_id']}` / `{item['key_group']}`: strict {strict['matched']}/{strict['total']}; normalised {normalised['matched']}/{normalised['total']} — indicative; measures keyword presence only. Missing strict: {', '.join(strict['missing_tokens']) or 'none'}; missing normalised: {', '.join(normalised['missing_tokens']) or 'none'}.")
        lines.extend(["", "### Other defects", ""])
        other = [d for d in metrics["defects"] if d["category"] not in {"text", "tables", "headings", "images", "contents", "mapping", "image_transcription"}]
        if other:
            for defect in other:
                lines.append(f"- {defect['category']} / {defect['code']}: {defect['message']}")
        else:
            lines.append("None.")
    return "\n".join(lines) + "\n"


def verify_all() -> dict[str, Any]:
    """Verify both local DOCX outputs and write machine-readable and readable reports."""
    results = []
    for doc_id in ("waterfall", "kanban"):
        metrics = verify_document(doc_id)
        results.append({"doc_id": doc_id, "metrics": metrics})
    overall = _aggregate(results)
    output = {"documents": {item["doc_id"]: item["metrics"] for item in results}, "overall": overall}
    output_dir = ROOT / "data/output"
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "metrics.json").write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n")
    (output_dir / "verification.md").write_text(_markdown_report(results, overall))
    return output
