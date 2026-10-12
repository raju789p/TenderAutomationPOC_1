"""Offline acceptance checks for the M4 Word renderer."""

import hashlib
import json
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from poc.render_docx import ROOT, render_document


DOCS = ("waterfall", "kanban")


def _all_document_text(document: Document) -> list[str]:
    texts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        texts.extend(cell.text for row in table.rows for cell in row.cells)
    for section in document.sections:
        texts.extend(paragraph.text for paragraph in section.footer.paragraphs)
        texts.extend(paragraph.text for paragraph in section.header.paragraphs)
    return texts


def _hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _display_hashes(doc_id: str) -> set[str]:
    elements = json.loads((ROOT / "data/work" / doc_id / "elements.json").read_text())["elements"]
    doc_dir = ROOT / "data/work" / doc_id
    return {_hash((doc_dir / item["image_path"]).read_bytes()) for item in elements if item["type"] == "figure"}


@pytest.mark.parametrize("doc_id", DOCS)
def test_docx_opens_and_all_elements_have_dispositions(doc_id: str, tmp_path: Path) -> None:
    ledger = render_document(doc_id, tmp_path / f"{doc_id}.docx")
    opened = Document(tmp_path / f"{doc_id}.docx")
    assert opened.paragraphs
    payload = json.loads((ROOT / "data/work" / doc_id / "elements.json").read_text())
    assert all(element["disposition"] in {"placed", "regenerated", "unplaced"} for element in payload["elements"])
    assert ledger["disposition_counts"]["pending"] == 0
    assert ledger["disposition_counts"]["unplaced"] == 0
    assert ledger["unplaced_ids"] == []


@pytest.mark.parametrize("doc_id", DOCS)
def test_source_tables_match_golden_fixtures(doc_id: str, tmp_path: Path) -> None:
    output = tmp_path / f"{doc_id}.docx"
    render_document(doc_id, output)
    elements = json.loads((ROOT / "data/work" / doc_id / "elements.json").read_text())["elements"]
    actual = [{"id": e["id"], "page": e["page"], "rows": e["rows"]} for e in elements if e["type"] == "table"]
    fixture_path = ROOT / "tests/fixtures" / f"golden_tables_{doc_id}.json"
    fixture = json.loads("\n".join(line for line in fixture_path.read_text().splitlines() if not line.startswith("//")))
    assert actual == fixture
    expected_rendered = [item["rows"] for item in fixture if next(e for e in elements if e["id"] == item["id"]).get("role") != "toc"]
    rendered_tables = [[[cell.text for cell in row.cells] for row in table.rows] for table in Document(output).tables[:len(expected_rendered)]]
    assert rendered_tables == expected_rendered
    appendix_headers = [cell.text for cell in Document(output).tables[-1].rows[0].cells]
    assert all("AI-generated" in heading for heading in appendix_headers)
    assert len(Document(output).tables[-1].rows) == sum(element["type"] == "figure" for element in elements) + 1


@pytest.mark.parametrize("doc_id", DOCS)
def test_docx_media_hashes_match_display_images(doc_id: str, tmp_path: Path) -> None:
    path = tmp_path / f"{doc_id}.docx"
    render_document(doc_id, path)
    with zipfile.ZipFile(path) as archive:
        media = {_hash(archive.read(name)) for name in archive.namelist() if name.startswith("word/media/")}
    assert media == _display_hashes(doc_id)


@pytest.mark.parametrize("doc_id", DOCS)
def test_render_is_deterministic(doc_id: str, tmp_path: Path) -> None:
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    render_document(doc_id, first)
    render_document(doc_id, second)
    assert _hash(first.read_bytes()) == _hash(second.read_bytes())


@pytest.mark.parametrize("doc_id", DOCS)
def test_rendered_text_has_no_private_use_glyphs(doc_id: str, tmp_path: Path) -> None:
    output = tmp_path / f"{doc_id}.docx"
    render_document(doc_id, output)
    document = Document(output)
    assert not any(0xE000 <= ord(character) <= 0xF8FF for text in _all_document_text(document) for character in text)


@pytest.mark.parametrize("doc_id", DOCS)
def test_table_pagination_and_leadin_keep_properties(doc_id: str, tmp_path: Path) -> None:
    output = tmp_path / f"{doc_id}.docx"
    render_document(doc_id, output)
    document = Document(output)
    assert document.tables
    for table in document.tables:
        for row in table.rows:
            tr_properties = row._tr.trPr
            assert tr_properties is not None
            assert tr_properties.find(qn("w:cantSplit")) is not None
        header_properties = table.rows[0]._tr.trPr
        assert header_properties is not None
        assert header_properties.find(qn("w:tblHeader")) is not None

    elements = json.loads((ROOT / "data/work" / doc_id / "elements.json").read_text())["elements"]
    label_texts = []
    after_table_texts = []
    for index, element in enumerate(elements[:-1]):
        following = next((candidate for candidate in elements[index + 1:] if candidate["type"] != "page_number"), None)
        if element["type"] == "paragraph" and following and following["type"] == "table" and following.get("role") != "toc":
            label_texts.append("".join(span["text"] for span in element.get("spans", [])))
        if element["type"] == "table" and element.get("role") != "toc" and following and following["type"] == "paragraph":
            after_table_texts.append("".join(span["text"] for span in following.get("spans", [])))
    for label_text in label_texts:
        rendered = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == label_text.strip())
        assert rendered.paragraph_format.keep_with_next
    for following_text in after_table_texts:
        rendered = next(paragraph for paragraph in document.paragraphs if paragraph.text.strip() == following_text.strip())
        assert rendered.paragraph_format.space_before.pt == pytest.approx(8)


@pytest.mark.parametrize("doc_id", DOCS)
def test_figure_paragraphs_stay_with_labels_and_captions(doc_id: str, tmp_path: Path) -> None:
    output = tmp_path / f"{doc_id}.docx"
    render_document(doc_id, output)
    document = Document(output)
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph._p.xpath(".//w:drawing"):
            assert paragraph.paragraph_format.keep_with_next
            label = paragraphs[index + 1]
            assert label.text.startswith("Figure ")
            assert label.paragraph_format.keep_with_next
            if index + 2 < len(paragraphs) and paragraphs[index + 2].text.startswith("AI-generated description:"):
                assert paragraphs[index + 2].paragraph_format.keep_together
