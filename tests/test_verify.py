"""Offline verifier checks, including deliberate DOCX corruption cases."""

import json
import re
import shutil
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from lxml import etree

from poc.render_docx import ROOT
from poc.verify import _source_text, _tokens, verify_document


def _copy_doc(doc_id: str, tmp_path: Path) -> Path:
    destination = tmp_path / f"{doc_id}.docx"
    shutil.copyfile(ROOT / "data/output" / f"{doc_id}_decision_brief.docx", destination)
    return destination


def _finding_codes(result: dict, category: str) -> list[str]:
    return [item["code"] for item in result[category]["defects"]]


def _source_paragraph(doc_id: str) -> tuple[dict, str, str]:
    elements = json.loads((ROOT / "data/work" / doc_id / "elements.json").read_text())["elements"]
    source_counts = Counter()
    for candidate in elements:
        if candidate.get("disposition") != "regenerated":
            source_counts.update(_tokens(_source_text(candidate)))
    for element in elements:
        if element["type"] == "paragraph" and element["role"] == "body" and element.get("disposition") == "placed" and len(_tokens(_source_text(element))) > 5:
            text = "".join(span["text"] for span in element.get("spans", []))
            unique = next((token for token in _tokens(text) if source_counts[token] == 1), None)
            if unique:
                return element, text, unique
    raise AssertionError("No suitable body paragraph found")


def _matching_paragraph(document: Document, source_text: str):
    wanted = " ".join(source_text.split())
    return next(paragraph for paragraph in document.paragraphs if " ".join(paragraph.text.split()) == wanted)


def test_clean_docx_outputs_verify_with_no_defects() -> None:
    for doc_id in ("waterfall", "kanban"):
        result = verify_document(doc_id)
        assert result["defects"] == []
        assert result["ledger"]["counts"]["unplaced"] == 0
        assert result["ledger"]["actual_regenerated_ids"] == result["ledger"]["expected_regenerated_ids"]
        assert result["text"]["missing_word_count"] == 0
        assert result["text"]["excess_word_count"] == 0
        assert result["images"]["set_equal"] and result["images"]["count_equal"]
        assert all(composite["matches"] for composite in result["images"]["composites"])


def test_delete_table_row_is_reported_as_table_shape_defect(tmp_path: Path) -> None:
    path = _copy_doc("waterfall", tmp_path)
    document = Document(path)
    table = document.tables[0]
    table._tbl.remove(table.rows[-1]._tr)
    document.save(path)
    result = verify_document("waterfall", path)
    assert _finding_codes(result, "tables") == ["table_shape_mismatch"]
    assert result["tables"]["defects"][0]["element_id"]


def test_swapped_images_are_reported_as_placement_defect(tmp_path: Path) -> None:
    path = _copy_doc("waterfall", tmp_path)
    namespaces = {
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }
    with zipfile.ZipFile(path) as source:
        files = {name: source.read(name) for name in source.namelist()}
    root = etree.fromstring(files["word/document.xml"])
    embeds = root.xpath(".//a:blip/@r:embed", namespaces=namespaces)
    assert len(embeds) >= 2 and embeds[0] != embeds[1]
    nodes = root.xpath(".//a:blip", namespaces=namespaces)
    first, second = nodes[0].get("{" + namespaces["r"] + "}embed"), nodes[1].get("{" + namespaces["r"] + "}embed")
    nodes[0].set("{" + namespaces["r"] + "}embed", second)
    nodes[1].set("{" + namespaces["r"] + "}embed", first)
    files["word/document.xml"] = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in files.items():
            target.writestr(name, data)
    result = verify_document("waterfall", path)
    assert result["images"]["set_equal"] and result["images"]["count_equal"]
    assert _finding_codes(result, "images") == ["image_placement_mismatch"]


def test_changed_paragraph_word_reports_one_missing_and_one_excess(tmp_path: Path) -> None:
    doc_id = "waterfall"
    path = _copy_doc(doc_id, tmp_path)
    element, source_text, token = _source_paragraph(doc_id)
    document = Document(path)
    paragraph = _matching_paragraph(document, source_text)
    match = re.search(re.escape(token), paragraph.text, flags=re.IGNORECASE)
    assert match
    paragraph.text = paragraph.text[:match.start()] + "CORRUPTIONWORD" + paragraph.text[match.end():]
    document.save(path)
    result = verify_document(doc_id, path)
    assert result["text"]["missing_word_count"] == 1
    assert result["text"]["excess_word_count"] == 1
    assert result["text"]["missing_words"] == [{"element_id": element["id"], "word": token.casefold(), "count": 1}]
    assert result["text"]["excess_words"][0]["word"] == "corruptionword"


def test_deleted_paragraph_reports_missing_words_at_source_element(tmp_path: Path) -> None:
    doc_id = "waterfall"
    path = _copy_doc(doc_id, tmp_path)
    element, source_text, unique_token = _source_paragraph(doc_id)
    document = Document(path)
    paragraph = _matching_paragraph(document, source_text)
    paragraph._p.getparent().remove(paragraph._p)
    document.save(path)
    result = verify_document(doc_id, path)
    assert {"element_id": element["id"], "word": unique_token, "count": 1} in result["text"]["missing_words"]


def test_changed_heading_title_reports_heading_mismatch(tmp_path: Path) -> None:
    path = _copy_doc("waterfall", tmp_path)
    document = Document(path)
    paragraph = next(p for p in document.paragraphs if p.style.name == "Heading 2" and p.text == "Abstract")
    paragraph.text = "Abstract changed"
    document.save(path)
    result = verify_document("waterfall", path)
    assert _finding_codes(result, "headings") == ["heading_missing_or_wrong_level"]
    assert result["headings"]["defects"][0]["element_id"] == "waterfall-p003-001"


def test_section_heading_moved_out_of_order_is_reported(tmp_path: Path) -> None:
    path = _copy_doc("kanban", tmp_path)
    document = Document(path)
    overview = next(p for p in document.paragraphs if p.style.name == "Heading 1" and p.text == "A. Overview")
    core = next(p for p in document.paragraphs if p.style.name == "Heading 1" and p.text == "B. Core Concept")
    overview._p.addprevious(core._p)
    document.save(path)
    result = verify_document("kanban", path)
    assert "target_section_order" in _finding_codes(result, "order")
    assert result["order"]["actual_section_order"][:2] == ["B_core_concept", "A_overview"]


def test_private_use_character_is_reported(tmp_path: Path) -> None:
    path = _copy_doc("waterfall", tmp_path)
    document = Document(path)
    document.paragraphs[0].add_run("\ue000")
    document.save(path)
    result = verify_document("waterfall", path)
    assert result["private_use"]["count"] == 1
    assert _finding_codes(result, "private_use") == ["private_use_character"]
